#!/usr/bin/env python3
"""Loop control helper: deterministic bookkeeping so agents spend tokens on work, not accounting.

  loopctl.py status  <loop>                 print loop status
  loopctl.py ready   <loop>                 exit 0 = work available, 1 = waiting on other loops, 2 = deadlocked on a blocked dep
  loopctl.py record  <loop> --start T --end T --runner-iter N --prompt-file F   (claude JSON on stdin)
  loopctl.py metrics <loop>                 regenerate the metrics table in progress.md
  loopctl.py summary                        metrics of all loops (stdout, markdown)
  loopctl.py excel                          build run-artifacts/prompts-and-sessions.xlsx (needs openpyxl)
"""
import argparse, csv, json, sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
LOOPS = ROOT / "loops"
ARTIFACTS = ROOT / "run-artifacts"
SESSIONS_CSV = ARTIFACTS / "sessions.csv"
MANUAL_CSV = ARTIFACTS / "manual-prompts.csv"
CSV_FIELDS = ["timestamp_start", "timestamp_end", "loop", "runner_iteration", "agent_iteration", "milestone",
              "mode", "session_id", "prompt", "result_summary", "is_error", "num_turns", "duration_s",
              "input_tokens", "output_tokens", "cache_read_tokens", "cache_creation_tokens", "total_tokens", "cost_usd"]
M_START, M_END = "<!-- METRICS:START -->", "<!-- METRICS:END -->"


def loop_dir(name):
    d = LOOPS / name
    if not d.is_dir():
        sys.exit(f"unknown loop: {name}")
    return d


def load_state(name):
    try:
        return json.loads((loop_dir(name) / "state" / "state.json").read_text())
    except json.JSONDecodeError as e:
        sys.exit(f"{name}: state.json is not valid JSON ({e})")


def load_iterations(name):
    f = loop_dir(name) / "state" / "iterations.jsonl"
    if not f.exists():
        return []
    return [json.loads(l) for l in f.read_text().splitlines() if l.strip()]


def cmd_status(a):
    print(load_state(a.loop).get("status", "unknown"))


def dep_status(dep, own_loop, own_state):
    loop, mid = dep.split(":", 1) if ":" in dep else (own_loop, dep)
    st = own_state if loop == own_loop else load_state(loop)
    for m in st.get("milestones", []):
        if m.get("id") == mid:
            return m.get("status")
    # the other loop has not planned yet -> not satisfiable yet
    return "pending"


def cmd_ready(a):
    st = load_state(a.loop)
    ms = st.get("milestones", [])
    if st.get("status") in ("not_started",) or not ms:
        sys.exit(0)  # planning is always possible
    if any(m.get("status") in ("in_progress", "verifying") for m in ms):
        sys.exit(0)
    pending = [m for m in ms if m.get("status") == "pending"]
    if not pending:
        sys.exit(0)  # all done/blocked -> agent runs FINAL or marks blocked
    deadlocked = True
    for m in pending:
        statuses = [dep_status(d, a.loop, st) for d in m.get("depends_on", [])]
        if all(s == "done" for s in statuses):
            sys.exit(0)
        if not any(s == "blocked" for s in statuses):
            deadlocked = False
    sys.exit(2 if deadlocked else 1)


def cmd_record(a):
    raw = sys.stdin.read()
    try:
        res = json.loads(raw)
    except json.JSONDecodeError:
        res = {"is_error": True, "result": raw[-500:]}
    st = load_state(a.loop)
    last = st.get("last_iteration") or {}
    u = res.get("usage") or {}
    it, ot = u.get("input_tokens", 0) or 0, u.get("output_tokens", 0) or 0
    cr, cc = u.get("cache_read_input_tokens", 0) or 0, u.get("cache_creation_input_tokens", 0) or 0
    row = {
        "timestamp_start": a.start, "timestamp_end": a.end, "loop": a.loop,
        "runner_iteration": a.runner_iter, "agent_iteration": st.get("iteration"),
        "milestone": last.get("milestone") or "UNATTRIBUTED", "mode": last.get("mode") or "",
        "session_id": res.get("session_id", ""),
        "prompt": Path(a.prompt_file).read_text() if a.prompt_file else "",
        "result_summary": (last.get("summary") or str(res.get("result", ""))[:300]).replace("\n", " "),
        "is_error": bool(res.get("is_error")), "num_turns": res.get("num_turns", 0),
        "duration_s": round((res.get("duration_ms") or 0) / 1000, 1),
        "input_tokens": it, "output_tokens": ot, "cache_read_tokens": cr, "cache_creation_tokens": cc,
        "total_tokens": it + ot + cr + cc, "cost_usd": round(res.get("total_cost_usd") or 0, 4),
    }
    with open(loop_dir(a.loop) / "state" / "iterations.jsonl", "a") as f:
        f.write(json.dumps(row) + "\n")
    ARTIFACTS.mkdir(exist_ok=True)
    new = not SESSIONS_CSV.exists()
    with open(SESSIONS_CSV, "a", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        if new:
            w.writeheader()
        w.writerow(row)
    write_metrics(a.loop)
    print(f"[{a.loop}] iter {a.runner_iter} milestone={row['milestone']} session={row['session_id']} "
          f"tokens={row['total_tokens']} cost=${row['cost_usd']} error={row['is_error']}")


def milestone_rows(name):
    st, its = load_state(name), load_iterations(name)
    order = ["PLANNING"] + [m["id"] for m in st.get("milestones", [])] + ["FINAL"]
    meta = {m["id"]: m for m in st.get("milestones", [])}
    groups = OrderedDict((k, []) for k in order)
    for r in its:
        groups.setdefault(r["milestone"], []).append(r)
    rows = []
    for mid, rs in groups.items():
        if not rs and mid in ("PLANNING", "FINAL"):
            continue
        m = meta.get(mid, {})
        status = m.get("status", "done" if rs else "")
        rows.append({
            "milestone": mid, "title": m.get("title", mid.title()), "status": status,
            "start": min((r["timestamp_start"] for r in rs), default=""),
            "end": max((r["timestamp_end"] for r in rs), default="") if status in ("done", "blocked") else "",
            "iterations": len(rs), "trials": m.get("trials", ""),
            "input": sum(r["input_tokens"] for r in rs), "output": sum(r["output_tokens"] for r in rs),
            "cache": sum(r["cache_read_tokens"] + r["cache_creation_tokens"] for r in rs),
            "total": sum(r["total_tokens"] for r in rs), "cost": sum(r["cost_usd"] for r in rs),
            "sessions": ", ".join(r["session_id"][:8] for r in rs if r["session_id"]),
        })
    return rows


def fmt_t(t):
    try:
        return datetime.fromisoformat(t).strftime("%Y-%m-%d %H:%M")
    except (ValueError, TypeError):
        return t or "-"


def metrics_md(name):
    rows = milestone_rows(name)
    if not rows:
        return "_No iterations yet._"
    out = ["| Milestone | Title | Status | Start | End | Iter | Trials | Input tok | Output tok | Cache tok | Total tok | Cost $ | Sessions |",
           "|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for r in rows:
        out.append(f"| {r['milestone']} | {r['title']} | {r['status']} | {fmt_t(r['start'])} | {fmt_t(r['end'])} | "
                   f"{r['iterations']} | {r['trials']} | {r['input']:,} | {r['output']:,} | {r['cache']:,} | "
                   f"{r['total']:,} | {r['cost']:.2f} | {r['sessions']} |")
    tot = lambda k: sum(r[k] for r in rows)
    out.append(f"| **TOTAL** | | | | | {tot('iterations')} | | {tot('input'):,} | {tot('output'):,} | "
               f"{tot('cache'):,} | {tot('total'):,} | {tot('cost'):.2f} | |")
    return "\n".join(out)


def write_metrics(name):
    p = loop_dir(name) / "progress.md"
    text = p.read_text()
    if M_START not in text or M_END not in text:
        text += f"\n\n## Milestone metrics\n{M_START}\n{M_END}\n"
    head, rest = text.split(M_START, 1)
    _, tail = rest.split(M_END, 1)
    p.write_text(f"{head}{M_START}\n{metrics_md(name)}\n{M_END}{tail}")


def cmd_metrics(a):
    write_metrics(a.loop)
    print(metrics_md(a.loop))


def cmd_summary(_):
    for d in sorted(p.name for p in LOOPS.iterdir() if (p / "state" / "state.json").exists()):
        print(f"### {d} ({load_state(d).get('status')})\n\n{metrics_md(d)}\n")


def cmd_excel(_):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        sys.exit("openpyxl missing: pip install --user openpyxl   (CSV data is already in run-artifacts/)")
    wb = Workbook()

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) for c in col[:50])))
        ws.freeze_panes = "A2"

    def read_csv(p):
        if not p.exists():
            return [], []
        with open(p, newline="") as f:
            rows = list(csv.reader(f))
        return (rows[0], rows[1:]) if rows else ([], [])

    h, rows = read_csv(SESSIONS_CSV)
    sheet("Loop Sessions", h or CSV_FIELDS, rows, first=True)
    h, rows = read_csv(MANUAL_CSV)
    sheet("Manual Sessions", h or ["date", "session_id", "tool", "purpose", "prompt"], rows)
    summ = []
    for d in sorted(p.name for p in LOOPS.iterdir() if (p / "state" / "state.json").exists()):
        for r in milestone_rows(d):
            summ.append([d, r["milestone"], r["title"], r["status"], r["start"], r["end"], r["iterations"],
                         r["trials"], r["input"], r["output"], r["cache"], r["total"], round(r["cost"], 4), r["sessions"]])
    sheet("Milestone Summary", ["loop", "milestone", "title", "status", "start", "end", "iterations", "trials",
                                "input_tokens", "output_tokens", "cache_tokens", "total_tokens", "cost_usd", "sessions"], summ)
    out = ARTIFACTS / "prompts-and-sessions.xlsx"
    wb.save(out)
    print(f"wrote {out}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = p.add_subparsers(dest="cmd", required=True)
    for c in ("status", "ready", "metrics"):
        sp.add_parser(c).add_argument("loop")
    r = sp.add_parser("record")
    r.add_argument("loop")
    r.add_argument("--start", required=True)
    r.add_argument("--end", required=True)
    r.add_argument("--runner-iter", type=int, default=0)
    r.add_argument("--prompt-file")
    sp.add_parser("summary")
    sp.add_parser("excel")
    a = p.parse_args()
    {"status": cmd_status, "ready": cmd_ready, "record": cmd_record, "metrics": cmd_metrics,
     "summary": cmd_summary, "excel": cmd_excel}[a.cmd](a)


if __name__ == "__main__":
    main()
